"""Unit tests for BR suppliers sales report and Excel generation"""
import io
from datetime import date
import openpyxl
from api.src.reports import export_service


def test_export_sales_by_br_suppliers_structure_and_sheets():
    mock_data = [
        {
            "sale_id": "11111111-1111-1111-1111-111111111111",
            "fecha_hora": "15/08/2026 10:30",
            "fecha_emision": date(2026, 8, 15),
            "comprobante": "001-001-0001234",
            "estado": "confirmado",
            "cliente_nombre": "JUAN PEREZ",
            "cliente_documento": "1234567-8",
            "supplier_id": "22222222-2222-2222-2222-222222222222",
            "proveedor_nombre": "JBS AVES DO BRASIL LTDA",
            "proveedor_ruc": "02.916.265/0001-60",
            "proveedor_pais": "Brasil",
            "tipo_proveedor": "brasilero",
            "product_id": "33333333-3333-3333-3333-333333333333",
            "sku": "BR-POLLO-001",
            "codigo_barra": "7891000100012",
            "producto_descripcion": "POLLO ENTERO CONGELADO BR 1KG",
            "cantidad": 5.0,
            "precio_unitario": 22000,
            "total_item": 110000,
            "medios_pago": "EFECTIVO",
        },
        {
            "sale_id": "11111111-1111-1111-1111-111111111112",
            "fecha_hora": "20/09/2026 14:15",
            "fecha_emision": date(2026, 9, 20),
            "comprobante": "001-001-0005678",
            "estado": "confirmado",
            "cliente_nombre": "MARIA GOMEZ",
            "cliente_documento": "8765432-1",
            "supplier_id": "44444444-4444-4444-4444-444444444444",
            "proveedor_nombre": "SEARA ALIMENTOS BR",
            "proveedor_ruc": "02.916.265/0002-40",
            "proveedor_pais": "Brasil",
            "tipo_proveedor": "brasilero",
            "product_id": "55555555-5555-5555-5555-555555555555",
            "sku": "BR-CERDO-002",
            "codigo_barra": "7891000200025",
            "producto_descripcion": "COSTILLA DE CERDO BR CONGELADA",
            "cantidad": 2.5,
            "precio_unitario": 40000,
            "total_item": 100000,
            "medios_pago": "QR DINELCO + PIX",
        },
    ]

    xlsx_bytes = export_service.export_sales_by_br_suppliers(
        data=mock_data,
        fecha_desde=date(2026, 8, 1),
        fecha_hasta=date(2026, 9, 30),
    )

    assert isinstance(xlsx_bytes, bytes)
    assert len(xlsx_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))

    # Comprobar hojas creadas (Hoja 1 = Agosto, Hoja 2 = Setiembre)
    sheet_names = wb.sheetnames
    assert "Agosto 2026" in sheet_names
    assert "Setiembre 2026" in sheet_names
    assert "Resumen por Proveedor" in sheet_names
    assert "Resumen Medios de Pago" in sheet_names

    # Validar hoja de Agosto
    ws_ago = wb["Agosto 2026"]
    assert "Agosto 2026" in str(ws_ago.cell(row=1, column=1).value)
    assert ws_ago.cell(row=3, column=4).value == "Proveedor"
    assert ws_ago.cell(row=3, column=5).value == "Origen / País"
    assert ws_ago.cell(row=4, column=4).value == "JBS AVES DO BRASIL LTDA"
    assert ws_ago.cell(row=4, column=15).value == "EFECTIVO"
    # Total Agosto
    tot_row_ago = 5
    assert ws_ago.cell(row=tot_row_ago, column=9).value == "TOTAL GENERAL"
    assert ws_ago.cell(row=tot_row_ago, column=10).value == 5.0
    assert ws_ago.cell(row=tot_row_ago, column=12).value == 110000.0

    # Validar hoja de Setiembre
    ws_set = wb["Setiembre 2026"]
    assert "Setiembre 2026" in str(ws_set.cell(row=1, column=1).value)
    assert ws_set.cell(row=4, column=4).value == "SEARA ALIMENTOS BR"
    assert ws_set.cell(row=4, column=15).value == "QR DINELCO + PIX"
    # Total Setiembre
    tot_row_set = 5
    assert ws_set.cell(row=tot_row_set, column=9).value == "TOTAL GENERAL"
    assert ws_set.cell(row=tot_row_set, column=10).value == 2.5
    assert ws_set.cell(row=tot_row_set, column=12).value == 100000.0

    # Validar hoja de proveedores
    ws2 = wb["Resumen por Proveedor"]
    assert "Resumen Consolidado por Proveedor" in str(ws2.cell(row=1, column=1).value)
    assert ws2.cell(row=4, column=2).value == "JBS AVES DO BRASIL LTDA"
    assert ws2.cell(row=5, column=2).value == "SEARA ALIMENTOS BR"

    # Validar hoja de medios de pago
    ws3 = wb["Resumen Medios de Pago"]
    assert "Participación por Medio de Pago" in str(ws3.cell(row=1, column=1).value)
    # Medios ordenados
    medios = [ws3.cell(row=r, column=2).value for r in range(4, 4 + len(mock_data))]
    assert "EFECTIVO" in medios
    assert "QR DINELCO + PIX" in medios


if __name__ == "__main__":
    test_export_sales_by_br_suppliers_structure_and_sheets()
    print("✓ All assertions passed in test_export_sales_by_br_suppliers_structure_and_sheets!")
