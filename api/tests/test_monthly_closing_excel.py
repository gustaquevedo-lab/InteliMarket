"""Unit tests for monthly closing Excel generator"""
import pytest
import io
import openpyxl
from api.src.financial.monthly_closing_excel import generate_closing_excel

def test_generate_closing_excel_multi_sheets_and_formulas():
    mock_data = {
        "periodo": {"mes": 9, "anio": 2026, "mes_nombre": "Septiembre"},
        "empresa": {
            "razon_social": "GRUPO SANTA TERESA E.A.S.",
            "nombre_fantasia": "Extra Supermercado Mayorista",
            "ruc": "80150377-9"
        },
        "ingresos": {
            "ventas_efectivo": 8682553757,
            "ventas_tarjetas_pos": 3885441242,
            "ventas_otros_medios": 0,
            "ventas_brutas": 12567994999,
            "descuentos": 0,
            "ventas_netas": 12567994999
        },
        "cogs": {
            "cmv_mercaderias": 10305755900,
            "mermas_frio_carniceria": 75395000,
            "total_cmv": 10381150900
        },
        "nomina": {
            "salario_base": 420000000,
            "horas_extras": 38500000,
            "masa_salarial_bruta": 458500000,
            "ips_patronal": 75652500,
            "aguinaldo_provisional": 38208333,
            "gastos_directos_personal": 28450000,
            "total_costo_laboral": 600810833,
            "empleados_activos": 84
        },
        "opex": {
            "ocupacion_servicios": {"total": 145800000},
            "mantenimiento_infraestructura": {"total": 54200000},
            "logistica_operacion": {"total": 98750000},
            "administracion_tecnologia": {"total": 62100000},
            "total_opex_no_laboral": 360850000
        },
        "financieros": {
            "comisiones_procesadoras_tarjetas": 85479000,
            "gastos_comisiones_bancarias": 14250000,
            "impuestos_y_tasas": 45000000,
            "total": 144729000
        },
        "ebitda": 1225183266,
        "resultado_neto": 1080454266,
        "conciliacion_bancaria": {
            "saldo_calibrado_oficial": 73874004,
            "cuentas_bancarias": [
                {"banco": "Banco Continental", "cuenta": "382364779308", "tipo": "Cta. Cte.", "saldo": 50816460},
                {"banco": "Banco Continental", "cuenta": "38250039706001", "tipo": "Caja Ahorro Dinelco", "saldo": 1674199},
                {"banco": "Banco Continental", "cuenta": "38250039706208", "tipo": "Ahorro PlugPay PIX", "saldo": 4736671},
                {"banco": "Banco Interfisa", "cuenta": "259058910", "tipo": "Cta. Cte.", "saldo": 9140355},
                {"banco": "Banco Interfisa", "cuenta": "259026699", "tipo": "Caja Ahorro PYG", "saldo": 7506319}
            ]
        }
    }
    
    excel_bytes = generate_closing_excel(mock_data)
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 1000
    
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=False)
    
    # Must have 8 sheets
    expected_sheets = [
        "1. Resumen Ejecutivo P&L",
        "2. Detalle Ingresos",
        "3. Detalle CMV",
        "4. Detalle Nomina SueldOK",
        "5. Ocupacion y Servicios",
        "6. Mantenimiento",
        "7. Logistica e Insumos",
        "8. Administracion y Finanzas"
    ]
    for s in expected_sheets:
        assert s in wb.sheetnames, f"Falta la hoja: {s}"
        
    ws_resumen = wb["1. Resumen Ejecutivo P&L"]
    # Check that live formulas are present in Resumen
    formulas_found = []
    for row in ws_resumen.iter_rows(values_only=False):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formulas_found.append(cell.value)
                
    assert len(formulas_found) >= 3, "Resumen Ejecutivo debe tener fórmulas vivas"
