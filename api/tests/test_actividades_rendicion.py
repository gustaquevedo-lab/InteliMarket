from datetime import date, datetime, timezone
from decimal import Decimal
import io
import uuid
import pytest
from unittest.mock import AsyncMock, MagicMock, patch
from fastapi import HTTPException
import openpyxl

from api.src.actividades.models import Actividad, ActividadTramo, ActividadAplicacion, ActividadRendicion
from api.src.actividades.rendicion import (
    obtener_previa_rendicion,
    cerrar_rendicion,
    generar_excel_rendicion,
)
from api.src.financial.models import SupplierCreditNote
from api.src.financial.schemas import SupplierCreditNoteCreate
from api.src.purchases.models import Supplier


@pytest.mark.asyncio
async def test_obtener_previa_rendicion_calcula_totales_correctamente():
    company_id = uuid.uuid4()
    supplier_id = uuid.uuid4()
    ciclo = "202610"

    act_escala = Actividad(
        id=uuid.uuid4(),
        company_id=company_id,
        supplier_id=supplier_id,
        nombre="Promo Escala 10%",
        tipo="escala",
        ciclo=ciclo,
        presupuesto=Decimal("500000"),
        consumido=Decimal("80000"),
    )
    act_bonif = Actividad(
        id=uuid.uuid4(),
        company_id=company_id,
        supplier_id=supplier_id,
        nombre="Promo Bonif 10+1",
        tipo="bonificacion",
        ciclo=ciclo,
        presupuesto=Decimal("300000"),
        consumido=Decimal("50000"),
    )

    db = AsyncMock()

    # 1. existing_rendicion -> None
    res_existing = MagicMock()
    res_existing.scalar_one_or_none.return_value = None

    # 2. acts -> [act_escala, act_bonif]
    res_acts = MagicMock()
    res_acts.scalars.return_value.all.return_value = [act_escala, act_bonif]

    # 3. apps group_by
    row_escala = MagicMock(actividad_id=act_escala.id, cant=2, descuentos=Decimal("80000"), total_valor=Decimal("80000"))
    row_bonif = MagicMock(actividad_id=act_bonif.id, cant=1, descuentos=Decimal("0"), total_valor=Decimal("50000"))
    res_apps = MagicMock()
    res_apps.all.return_value = [row_escala, row_bonif]

    # 4. supplier
    sup = Supplier(id=supplier_id, razon_social="ARCOR PARAGUAY S.A.")
    res_sup = MagicMock()
    res_sup.scalar_one_or_none.return_value = sup

    db.execute.side_effect = [res_existing, res_acts, res_apps, res_sup]

    previa = await obtener_previa_rendicion(db, company_id, supplier_id, ciclo)

    assert previa["supplier_id"] == str(supplier_id)
    assert previa["supplier_nombre"] == "ARCOR PARAGUAY S.A."
    assert previa["ciclo"] == ciclo
    assert previa["puede_cerrar"] is True
    assert previa["total_descuentos"] == 80000.0
    assert previa["total_bonificaciones"] == 50000.0
    assert previa["total_rendicion"] == 130000.0
    assert previa["cantidad_aplicaciones"] == 3
    assert len(previa["actividades"]) == 2


@pytest.mark.asyncio
async def test_cerrar_rendicion_crea_asiento_y_vincula_aplicaciones():
    company_id = uuid.uuid4()
    supplier_id = uuid.uuid4()
    user_id = uuid.uuid4()
    ciclo = "202610"

    act = Actividad(
        id=uuid.uuid4(),
        company_id=company_id,
        supplier_id=supplier_id,
        nombre="Promo Escala 5%",
        tipo="escala",
        ciclo=ciclo,
        presupuesto=Decimal("200000"),
        consumido=Decimal("60000"),
        cuenta_contable="1152",
    )

    app1 = ActividadAplicacion(
        id=uuid.uuid4(),
        company_id=company_id,
        actividad_id=act.id,
        sale_id=uuid.uuid4(),
        descuento_monto=Decimal("35000"),
        valor=Decimal("35000"),
        rendicion_id=None,
    )
    app2 = ActividadAplicacion(
        id=uuid.uuid4(),
        company_id=company_id,
        actividad_id=act.id,
        sale_id=uuid.uuid4(),
        descuento_monto=Decimal("25000"),
        valor=Decimal("25000"),
        rendicion_id=None,
    )

    db = AsyncMock()

    # 1. existing_rendicion -> None
    res_existing = MagicMock()
    res_existing.scalar_one_or_none.return_value = None

    # 2. acts
    res_acts = MagicMock()
    res_acts.scalars.return_value.all.return_value = [act]

    # 3. apps
    res_apps = MagicMock()
    res_apps.scalars.return_value.all.return_value = [app1, app2]

    # 4. cuenta_debe
    acc_debe = MagicMock(id=uuid.uuid4(), codigo="1152")
    res_acc_debe = MagicMock()
    res_acc_debe.scalar_one_or_none.return_value = acc_debe

    # 5. cuenta_haber
    acc_haber = MagicMock(id=uuid.uuid4(), codigo="4110")
    res_acc_haber = MagicMock()
    res_acc_haber.scalar_one_or_none.return_value = acc_haber

    # 6. period
    period = MagicMock(id=uuid.uuid4())

    # 7. supplier
    sup = Supplier(id=supplier_id, razon_social="ARCOR PARAGUAY S.A.")
    res_sup = MagicMock()
    res_sup.scalar_one_or_none.return_value = sup

    db.execute.side_effect = [
        res_existing,
        res_acts,
        res_apps,
        res_acc_debe,
        res_acc_haber,
        res_sup,
    ]

    with patch("api.src.actividades.rendicion._resolve_period_for_manual", return_value=period), \
         patch("api.src.actividades.rendicion.post_accounting_entry", new=AsyncMock(return_value={"entry": MagicMock(id=uuid.uuid4())})) as mock_post:

        res = await cerrar_rendicion(db, company_id, supplier_id, ciclo, user_id)

        assert res["success"] is True
        assert res["total_rendicion"] == 60000.0
        assert res["aplicaciones_vinculadas"] == 2
        assert res["estado"] == "cerrada"
        assert app1.rendicion_id is not None
        assert app2.rendicion_id == app1.rendicion_id
        # Se generaron 2 partidas contables (Debe y Haber)
        assert mock_post.call_count == 2


@pytest.mark.asyncio
async def test_cerrar_rendicion_duplicada_lanza_409():
    company_id = uuid.uuid4()
    supplier_id = uuid.uuid4()
    ciclo = "202610"

    existing_rend = ActividadRendicion(
        id=uuid.uuid4(),
        company_id=company_id,
        supplier_id=supplier_id,
        ciclo=ciclo,
        estado="cerrada",
    )

    db = AsyncMock()
    res_existing = MagicMock()
    res_existing.scalar_one_or_none.return_value = existing_rend
    db.execute.return_value = res_existing

    with pytest.raises(HTTPException) as exc:
        await cerrar_rendicion(db, company_id, supplier_id, ciclo, uuid.uuid4())
    assert exc.value.status_code == 409


@pytest.mark.asyncio
async def test_nota_credito_proveedor_concilia_rendicion():
    from api.src.financial.service import create_supplier_credit_note

    company_id = uuid.uuid4()
    supplier_id = uuid.uuid4()
    rendicion_id = uuid.uuid4()

    rendicion = ActividadRendicion(
        id=rendicion_id,
        company_id=company_id,
        supplier_id=supplier_id,
        ciclo="202610",
        total_rendicion=Decimal("100000"),
        monto=Decimal("100000"),
        estado="cerrada",
    )

    db = AsyncMock()

    # 1. find supplier
    sup = Supplier(id=supplier_id, company_id=company_id, razon_social="ARCOR")
    res_sup = MagicMock()
    res_sup.scalar_one_or_none.return_value = sup

    # 2. find rendicion
    res_rend = MagicMock()
    res_rend.scalar_one_or_none.return_value = rendicion

    # 3. prev_nc sum -> 0
    res_nc_sum = MagicMock()
    res_nc_sum.scalar.return_value = Decimal("0")

    # 4. _get_company_info query
    res_comp = MagicMock()
    res_comp.first.return_value = None

    db.execute.side_effect = [res_sup, res_rend, res_nc_sum, res_comp]

    data = SupplierCreditNoteCreate(
        supplier_id=str(supplier_id),
        numero="NC-001-001-000001",
        timbrado="12345678",
        fecha=date.today(),
        monto=100000.0,
        motivo="Liquidacion recupero actividades 202610",
        rendicion_id=str(rendicion_id),
    )

    note = await create_supplier_credit_note(db, str(company_id), data)
    assert note["rendicion_id"] == str(rendicion_id)
    assert rendicion.estado == "conciliada"


@pytest.mark.asyncio
async def test_generar_excel_rendicion_crea_resumen_y_detalle():
    company_id = uuid.uuid4()
    supplier_id = uuid.uuid4()
    rendicion_id = uuid.uuid4()

    rendicion = ActividadRendicion(
        id=rendicion_id,
        company_id=company_id,
        supplier_id=supplier_id,
        ciclo="202610",
        total_descuentos=Decimal("40000"),
        total_bonificaciones=Decimal("15000"),
        total_rendicion=Decimal("55000"),
        monto=Decimal("55000"),
        estado="cerrada",
        cerrada_at=datetime.now(timezone.utc),
    )

    sup = Supplier(id=supplier_id, razon_social="NESTLE PARAGUAY S.A.")

    row_sum = MagicMock(
        nombre="Promo Oblea 10%",
        tipo="escala",
        presupuesto=Decimal("200000"),
        consumido=Decimal("55000"),
        descuentos=Decimal("40000"),
        total_valor=Decimal("40000"),
    )

    row_det = MagicMock(
        fecha=datetime(2026, 10, 20, 10, 0, 0),
        numero_factura="001-001-0001234",
        cliente_ruc="80012345-6",
        cliente_nombre="Despensa San Cayetano",
        actividad_nombre="Promo Oblea 10%",
        actividad_tipo="escala",
        producto_sku="WAFER-CHOCO-40",
        producto_nombre="Oblea Chocolate 40u",
        item_cantidad=Decimal("2.0"),
        item_desc_monto=Decimal("40000"),
        descuento_monto=Decimal("40000"),
        unidades_bonificadas=Decimal("0"),
        valor=Decimal("40000"),
    )

    db = AsyncMock()

    # 1. rendicion
    res_rend = MagicMock()
    res_rend.scalar_one_or_none.return_value = rendicion
    # 2. supplier
    res_sup = MagicMock()
    res_sup.scalar_one_or_none.return_value = sup
    # 3. summary
    res_sum = MagicMock()
    res_sum.all.return_value = [row_sum]
    # 4. detail
    res_det = MagicMock()
    res_det.all.return_value = [row_det]

    db.execute.side_effect = [res_rend, res_sup, res_sum, res_det]

    excel_bytes = await generar_excel_rendicion(db, rendicion_id)
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Resumen" in wb.sheetnames
    assert "Detalle" in wb.sheetnames

    ws_resumen = wb["Resumen"]
    assert "RENDICIÓN DE ACTIVIDADES" in str(ws_resumen["A1"].value)
    assert "NESTLE PARAGUAY S.A." in str(ws_resumen["A2"].value)

    ws_detalle = wb["Detalle"]
    assert ws_detalle["A1"].value == "Fecha"
    assert ws_detalle["B1"].value == "Comprobante"
    assert ws_detalle["B2"].value == "001-001-0001234"
    assert ws_detalle["D2"].value == "Despensa San Cayetano"
