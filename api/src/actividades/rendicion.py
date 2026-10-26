"""Servicio de Rendición de Actividades al Proveedor:
- Vista previa de liquidación
- Cierre formal con asiento contable de partida doble
- Exportación a Excel profesional (Resumen + Detalle)
"""

from datetime import date, datetime, timezone
from decimal import Decimal
import io
import uuid
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select, func, and_
from sqlalchemy.ext.asyncio import AsyncSession
from fastapi import HTTPException

from api.src.actividades.models import Actividad, ActividadAplicacion, ActividadRendicion
from api.src.sales.models import Sale, SaleItem
from api.src.products.models import Product
from api.src.purchases.models import Supplier
from api.src.customers.models import Customer
from api.src.integrated_finance.models import AccountPlan, AccountingPeriod
from api.src.integrated_finance.schemas import AccountingEntryCreate
from api.src.integrated_finance.service import post_accounting_entry, _resolve_period_for_manual


async def obtener_previa_rendicion(
    db: AsyncSession,
    company_id: uuid.UUID | str,
    supplier_id: uuid.UUID | str,
    ciclo: str,
) -> dict:
    cid = uuid.UUID(str(company_id))
    sid = uuid.UUID(str(supplier_id))

    # Verificar si ya existe rendición cerrada o conciliada
    existing_stmt = select(ActividadRendicion).where(
        ActividadRendicion.company_id == cid,
        ActividadRendicion.supplier_id == sid,
        ActividadRendicion.ciclo == ciclo,
        ActividadRendicion.estado.in_(["cerrada", "conciliada"])
    )
    existing_res = await db.execute(existing_stmt)
    existing_rendicion = existing_res.scalar_one_or_none()

    # Obtener actividades del proveedor en el ciclo
    acts_stmt = select(Actividad).where(
        Actividad.company_id == cid,
        Actividad.supplier_id == sid,
        Actividad.ciclo == ciclo,
    ).order_by(Actividad.nombre)
    acts_res = await db.execute(acts_stmt)
    acts = acts_res.scalars().all()
    act_ids = [a.id for a in acts]

    total_descuentos = Decimal(0)
    total_bonificaciones = Decimal(0)
    cantidad_aplicaciones = 0
    act_summary = []

    if act_ids:
        # Aplicaciones pendientes de rendición (rendicion_id IS NULL)
        apps_stmt = select(
            ActividadAplicacion.actividad_id,
            func.count(ActividadAplicacion.id).label("cant"),
            func.sum(ActividadAplicacion.descuento_monto).label("descuentos"),
            func.sum(ActividadAplicacion.valor).label("total_valor"),
        ).where(
            ActividadAplicacion.company_id == cid,
            ActividadAplicacion.actividad_id.in_(act_ids),
            ActividadAplicacion.rendicion_id == None,
        ).group_by(ActividadAplicacion.actividad_id)

        apps_res = await db.execute(apps_stmt)
        app_data_by_act = {
            row.actividad_id: {
                "cant": row.cant,
                "descuentos": Decimal(str(row.descuentos or 0)),
                "total_valor": Decimal(str(row.total_valor or 0)),
            }
            for row in apps_res.all()
        }

        for act in acts:
            info = app_data_by_act.get(act.id, {"cant": 0, "descuentos": Decimal(0), "total_valor": Decimal(0)})
            desc = info["descuentos"]
            boni = max(Decimal(0), info["total_valor"] - desc) if act.tipo == "bonificacion" else Decimal(0)
            if act.tipo == "bonificacion":
                desc = Decimal(0)
                boni = info["total_valor"]

            total_descuentos += desc
            total_bonificaciones += boni
            cantidad_aplicaciones += info["cant"]

            act_summary.append({
                "actividad_id": str(act.id),
                "nombre": act.nombre,
                "tipo": act.tipo,
                "presupuesto": float(act.presupuesto or 0),
                "consumido": float(act.consumido or 0),
                "aplicaciones_count": info["cant"],
                "total_descuentos": float(desc),
                "total_bonificaciones": float(boni),
                "total_rendicion": float(desc + boni),
            })

    total_rendicion = total_descuentos + total_bonificaciones

    # Obtener nombre proveedor
    sup_res = await db.execute(select(Supplier).where(Supplier.id == sid))
    sup = sup_res.scalar_one_or_none()

    return {
        "supplier_id": str(sid),
        "supplier_nombre": sup.razon_social if sup else "",
        "ciclo": ciclo,
        "puede_cerrar": existing_rendicion is None and cantidad_aplicaciones > 0,
        "ya_rendida": existing_rendicion is not None,
        "rendicion_existente": {
            "id": str(existing_rendicion.id),
            "estado": existing_rendicion.estado,
            "total_rendicion": float(existing_rendicion.total_rendicion or existing_rendicion.monto or 0),
        } if existing_rendicion else None,
        "total_descuentos": float(total_descuentos),
        "total_bonificaciones": float(total_bonificaciones),
        "total_rendicion": float(total_rendicion),
        "cantidad_aplicaciones": cantidad_aplicaciones,
        "actividades": act_summary,
    }


async def cerrar_rendicion(
    db: AsyncSession,
    company_id: uuid.UUID | str,
    supplier_id: uuid.UUID | str,
    ciclo: str,
    user_id: uuid.UUID | str,
) -> dict:
    cid = uuid.UUID(str(company_id))
    sid = uuid.UUID(str(supplier_id))
    uid = uuid.UUID(str(user_id)) if user_id else None

    # 1. Validar que no haya rendición cerrada previa para este ciclo y proveedor
    existing_stmt = select(ActividadRendicion).where(
        ActividadRendicion.company_id == cid,
        ActividadRendicion.supplier_id == sid,
        ActividadRendicion.ciclo == ciclo,
        ActividadRendicion.estado.in_(["cerrada", "conciliada"])
    )
    existing_res = await db.execute(existing_stmt)
    if existing_res.scalar_one_or_none():
        raise HTTPException(status_code=409, detail=f"Ya existe una rendición cerrada para el proveedor y ciclo {ciclo}")

    # 2. Buscar actividades y aplicaciones huérfanas
    acts_stmt = select(Actividad).where(
        Actividad.company_id == cid,
        Actividad.supplier_id == sid,
        Actividad.ciclo == ciclo,
    )
    acts_res = await db.execute(acts_stmt)
    acts = acts_res.scalars().all()
    act_ids = [a.id for a in acts]
    if not act_ids:
        raise HTTPException(status_code=400, detail="No hay actividades para este proveedor en el ciclo especificado")

    apps_stmt = select(ActividadAplicacion).where(
        ActividadAplicacion.company_id == cid,
        ActividadAplicacion.actividad_id.in_(act_ids),
        ActividadAplicacion.rendicion_id == None,
    )
    apps_res = await db.execute(apps_stmt)
    apps = apps_res.scalars().all()
    if not apps:
        raise HTTPException(status_code=400, detail="No hay aplicaciones pendientes de rendición")

    act_map = {a.id: a for a in acts}
    total_descuentos = Decimal(0)
    total_bonificaciones = Decimal(0)

    for app in apps:
        act = act_map.get(app.actividad_id)
        if act and act.tipo == "bonificacion":
            total_bonificaciones += Decimal(str(app.valor or 0))
        else:
            total_descuentos += Decimal(str(app.descuento_monto or 0))

    total_rendicion = total_descuentos + total_bonificaciones

    # 3. Crear ActividadRendicion
    rendicion_id = uuid.uuid4()
    rendicion = ActividadRendicion(
        id=rendicion_id,
        company_id=cid,
        supplier_id=sid,
        ciclo=ciclo,
        total_descuentos=total_descuentos,
        total_bonificaciones=total_bonificaciones,
        total_rendicion=total_rendicion,
        monto=total_rendicion,
        estado="cerrada",
        cerrada_por=uid,
        cerrada_at=datetime.now(timezone.utc),
        created_by=uid,
    )
    db.add(rendicion)
    await db.flush()

    # 4. Vincular aplicaciones con rendicion_id
    for app in apps:
        app.rendicion_id = rendicion.id

    # 5. Generar asiento contable de partida doble
    # Debe: Cuenta contable de la actividad o default "1152"
    # Haber: Cuenta de Ventas "4110" (o "31111" / "4.1.01")
    cuenta_debe_codigo = None
    for a in acts:
        if a.cuenta_contable:
            cuenta_debe_codigo = a.cuenta_contable.strip()
            break
    if not cuenta_debe_codigo:
        cuenta_debe_codigo = "1152"

    # Buscar cuentas contables en account_plans
    cuenta_debe = None
    for cod in [cuenta_debe_codigo, "1152", "11213", "1.1.02"]:
        res_acc = await db.execute(select(AccountPlan).where(AccountPlan.company_id == cid, AccountPlan.codigo == cod))
        cuenta_debe = res_acc.scalar_one_or_none()
        if cuenta_debe:
            break

    cuenta_haber = None
    for cod in ["4110", "31111", "4.1.01", "311"]:
        res_acc = await db.execute(select(AccountPlan).where(AccountPlan.company_id == cid, AccountPlan.codigo == cod))
        cuenta_haber = res_acc.scalar_one_or_none()
        if cuenta_haber:
            break

    hoy = date.today()
    period = await _resolve_period_for_manual(db, str(cid), hoy)

    sup_res = await db.execute(select(Supplier).where(Supplier.id == sid))
    sup = sup_res.scalar_one_or_none()
    sup_nombre = sup.razon_social if sup else str(sid)

    asiento_num = f"REND-{ciclo}-{rendicion.id.hex[:6].upper()}"
    concepto = f"Rendición actividades {ciclo} - Proveedor {sup_nombre}"

    entry_debe_id = None
    if cuenta_debe and cuenta_haber and total_rendicion > 0:
        res_debe = await post_accounting_entry(
            db,
            AccountingEntryCreate(
                company_id=str(cid),
                period_id=str(period.id),
                account_id=str(cuenta_debe.id),
                fecha=hoy,
                tipo="debe",
                monto=float(total_rendicion),
                concepto=concepto,
                referencia_tipo="actividad_rendicion",
                referencia_id=str(rendicion.id),
                asiento_numero=asiento_num,
            ),
            user_id=str(uid) if uid else None,
        )
        if res_debe.get("entry"):
            entry_debe_id = res_debe["entry"].id

        await post_accounting_entry(
            db,
            AccountingEntryCreate(
                company_id=str(cid),
                period_id=str(period.id),
                account_id=str(cuenta_haber.id),
                fecha=hoy,
                tipo="haber",
                monto=float(total_rendicion),
                concepto=concepto,
                referencia_tipo="actividad_rendicion",
                referencia_id=str(rendicion.id),
                asiento_numero=asiento_num,
            ),
            user_id=str(uid) if uid else None,
        )

    if entry_debe_id:
        rendicion.asiento_id = entry_debe_id

    await db.flush()
    await db.refresh(rendicion)

    return {
        "success": True,
        "rendicion_id": str(rendicion.id),
        "total_rendicion": float(total_rendicion),
        "total_descuentos": float(total_descuentos),
        "total_bonificaciones": float(total_bonificaciones),
        "aplicaciones_vinculadas": len(apps),
        "asiento_numero": asiento_num,
        "estado": rendicion.estado,
    }


async def generar_excel_rendicion(db: AsyncSession, rendicion_id: uuid.UUID | str) -> bytes:
    rid = uuid.UUID(str(rendicion_id))

    r_stmt = select(ActividadRendicion).where(ActividadRendicion.id == rid)
    r_res = await db.execute(r_stmt)
    rendicion = r_res.scalar_one_or_none()
    if not rendicion:
        raise HTTPException(status_code=404, detail="Rendición no encontrada")

    sup_res = await db.execute(select(Supplier).where(Supplier.id == rendicion.supplier_id))
    supplier = sup_res.scalar_one_or_none()
    sup_nombre = supplier.razon_social if supplier else "Proveedor"

    wb = openpyxl.Workbook()

    # ── ESTILOS ──
    header_fill = PatternFill(start_color="0F1F3D", end_color="0F1F3D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="0F1F3D")
    subtitle_font = Font(name="Calibri", size=11, bold=False, color="475569")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    money_format = '#,##0'
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    # ── HOJA 1: RESUMEN ──
    ws_resumen = wb.active
    ws_resumen.title = "Resumen"

    ws_resumen.cell(row=1, column=1, value=f"RENDICIÓN DE ACTIVIDADES DE PROVEEDOR").font = title_font
    ws_resumen.cell(row=2, column=1, value=f"Proveedor: {sup_nombre} | Ciclo: {rendicion.ciclo} | Estado: {rendicion.estado.upper()}").font = subtitle_font
    ws_resumen.cell(row=3, column=1, value=f"Fecha de cierre: {rendicion.cerrada_at.strftime('%Y-%m-%d %H:%M') if rendicion.cerrada_at else 'Pendiente'}").font = subtitle_font

    headers_resumen = [
        "Actividad",
        "Tipo",
        "Presupuesto (₲)",
        "Consumido (₲)",
        "Descuentos (₲)",
        "Bonificaciones (₲)",
        "Total Rendición (₲)"
    ]

    for col_idx, h in enumerate(headers_resumen, 1):
        cell = ws_resumen.cell(row=5, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Obtener actividades y agregados para esta rendición
    summary_stmt = select(
        Actividad.nombre,
        Actividad.tipo,
        Actividad.presupuesto,
        Actividad.consumido,
        func.sum(ActividadAplicacion.descuento_monto).label("descuentos"),
        func.sum(ActividadAplicacion.valor).label("total_valor"),
    ).join(
        ActividadAplicacion, ActividadAplicacion.actividad_id == Actividad.id
    ).where(
        ActividadAplicacion.rendicion_id == rid
    ).group_by(
        Actividad.id, Actividad.nombre, Actividad.tipo, Actividad.presupuesto, Actividad.consumido
    ).order_by(Actividad.nombre)

    summary_res = await db.execute(summary_stmt)
    rows_summary = summary_res.all()

    current_row = 6
    tot_presup = Decimal(0)
    tot_cons = Decimal(0)
    tot_desc = Decimal(0)
    tot_boni = Decimal(0)
    tot_rend = Decimal(0)

    for r in rows_summary:
        nom = r.nombre
        tipo = r.tipo
        presup = Decimal(str(r.presupuesto or 0))
        cons = Decimal(str(r.consumido or 0))
        desc = Decimal(str(r.descuentos or 0))
        boni = max(Decimal(0), Decimal(str(r.total_valor or 0)) - desc) if tipo == "bonificacion" else Decimal(0)
        if tipo == "bonificacion":
            desc = Decimal(0)
            boni = Decimal(str(r.total_valor or 0))
        rend = desc + boni

        tot_presup += presup
        tot_cons += cons
        tot_desc += desc
        tot_boni += boni
        tot_rend += rend

        ws_resumen.cell(row=current_row, column=1, value=nom).font = regular_font
        ws_resumen.cell(row=current_row, column=2, value=tipo.capitalize()).font = regular_font
        ws_resumen.cell(row=current_row, column=3, value=float(presup)).number_format = money_format
        ws_resumen.cell(row=current_row, column=4, value=float(cons)).number_format = money_format
        ws_resumen.cell(row=current_row, column=5, value=float(desc)).number_format = money_format
        ws_resumen.cell(row=current_row, column=6, value=float(boni)).number_format = money_format
        ws_resumen.cell(row=current_row, column=7, value=float(rend)).number_format = money_format

        for c in range(1, 8):
            ws_resumen.cell(row=current_row, column=c).border = thin_border
        current_row += 1

    # Fila de totales
    ws_resumen.cell(row=current_row, column=1, value="TOTAL GENERAL").font = bold_font
    ws_resumen.cell(row=current_row, column=2, value="").font = bold_font
    ws_resumen.cell(row=current_row, column=3, value=float(tot_presup)).number_format = money_format
    ws_resumen.cell(row=current_row, column=4, value=float(tot_cons)).number_format = money_format
    ws_resumen.cell(row=current_row, column=5, value=float(tot_desc)).number_format = money_format
    ws_resumen.cell(row=current_row, column=6, value=float(tot_boni)).number_format = money_format
    ws_resumen.cell(row=current_row, column=7, value=float(tot_rend)).number_format = money_format

    for c in range(1, 8):
        cell = ws_resumen.cell(row=current_row, column=c)
        cell.font = bold_font
        cell.border = thin_border
        cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    for col in ws_resumen.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_resumen.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # ── HOJA 2: DETALLE ──
    ws_detalle = wb.create_sheet(title="Detalle")

    headers_detalle = [
        "Fecha",
        "Comprobante",
        "RUC Cliente",
        "Cliente",
        "Actividad",
        "Tipo",
        "SKU",
        "Producto",
        "Cantidad",
        "Descuento Unit. (₲)",
        "Descuento Total (₲)",
        "Valor Bonificado (₲)",
    ]

    for col_idx, h in enumerate(headers_detalle, 1):
        cell = ws_detalle.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    detail_stmt = select(
        Sale.created_at.label("fecha"),
        Sale.numero.label("numero_factura"),
        Customer.ruc.label("cliente_ruc"),
        Customer.razon_social.label("cliente_nombre"),
        Actividad.nombre.label("actividad_nombre"),
        Actividad.tipo.label("actividad_tipo"),
        Product.sku.label("producto_sku"),
        Product.nombre.label("producto_nombre"),
        SaleItem.cantidad.label("item_cantidad"),
        SaleItem.descuento_monto.label("item_desc_monto"),
        ActividadAplicacion.descuento_monto,
        ActividadAplicacion.unidades_bonificadas,
        ActividadAplicacion.valor,
    ).select_from(
        ActividadAplicacion
    ).join(
        Actividad, Actividad.id == ActividadAplicacion.actividad_id
    ).join(
        Sale, Sale.id == ActividadAplicacion.sale_id
    ).outerjoin(
        Customer, Customer.id == Sale.customer_id
    ).outerjoin(
        SaleItem, SaleItem.id == ActividadAplicacion.sale_item_id
    ).outerjoin(
        Product, Product.id == SaleItem.product_id
    ).where(
        ActividadAplicacion.rendicion_id == rid
    ).order_by(
        Sale.created_at, Sale.numero
    )

    detail_res = await db.execute(detail_stmt)
    rows_detail = detail_res.all()

    det_row = 2
    for d in rows_detail:
        fec_str = d.fecha.strftime("%Y-%m-%d") if d.fecha else ""
        num_comp = d.numero_factura or "-"
        ruc = d.cliente_ruc or ""
        cli = d.cliente_nombre or "Cliente Final"
        act_nom = d.actividad_nombre
        tipo = d.actividad_tipo
        sku = d.producto_sku or ""
        prod = d.producto_nombre or ""
        cant = float(d.item_cantidad or 0)
        desc_tot = float(d.descuento_monto or 0)
        desc_u = round(desc_tot / cant) if cant > 0 else 0.0
        val_boni = float(d.valor or 0) if tipo == "bonificacion" else 0.0

        ws_detalle.cell(row=det_row, column=1, value=fec_str).font = regular_font
        ws_detalle.cell(row=det_row, column=2, value=num_comp).font = regular_font
        ws_detalle.cell(row=det_row, column=3, value=ruc).font = regular_font
        ws_detalle.cell(row=det_row, column=4, value=cli).font = regular_font
        ws_detalle.cell(row=det_row, column=5, value=act_nom).font = regular_font
        ws_detalle.cell(row=det_row, column=6, value=tipo.capitalize()).font = regular_font
        ws_detalle.cell(row=det_row, column=7, value=sku).font = regular_font
        ws_detalle.cell(row=det_row, column=8, value=prod).font = regular_font
        ws_detalle.cell(row=det_row, column=9, value=cant).number_format = '#,##0.00'
        ws_detalle.cell(row=det_row, column=10, value=desc_u).number_format = money_format
        ws_detalle.cell(row=det_row, column=11, value=desc_tot).number_format = money_format
        ws_detalle.cell(row=det_row, column=12, value=val_boni).number_format = money_format

        for c in range(1, 13):
            ws_detalle.cell(row=det_row, column=c).border = thin_border
        det_row += 1

    for col in ws_detalle.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_detalle.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
