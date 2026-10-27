import asyncio
import uuid
import datetime
from decimal import Decimal
import openpyxl
from sqlalchemy import text
from api.src.db import async_session_factory

GONZALITO_ID = 'a2f37cf3-caa5-42b5-84cd-270a549d5a73'
COMPANY_ID = '00000000-0000-0000-0000-000000000010'
TIMBRADO_DEF = '17090459'

async def run_remediation():
    wb = openpyxl.load_workbook('datos/FACTURAS PENDIENTES GONZALITO.xlsx', data_only=True)
    sheet = wb.active

    excel_invoices = []
    excel_ncs = []
    current_inv = None

    for i in range(1, sheet.max_row + 1):
        doc = sheet.cell(i, 2).value
        amt = sheet.cell(i, 3).value
        date_v = sheet.cell(i, 1).value
        if doc is None or amt is None:
            continue
        doc_str = str(doc).strip()
        amt_d = Decimal(str(amt))

        if amt_d > 0:
            current_inv = {
                'row': i,
                'numero_factura': doc_str,
                'fecha': date_v,
                'total_papel': amt_d,
                'ncs': []
            }
            excel_invoices.append(current_inv)
        else:
            nc_info = {
                'row': i,
                'numero_nc': doc_str,
                'fecha': date_v,
                'monto': abs(amt_d),
                'factura_origen': current_inv['numero_factura'] if current_inv else None
            }
            excel_ncs.append(nc_info)
            if current_inv:
                current_inv['ncs'].append(nc_info)

    print(f"Excel procesado: {len(excel_invoices)} facturas, {len(excel_ncs)} NCs asociadas")

    # Nuevas NCs recibidas hoy indicadas por el usuario
    new_ncs_today = [
        {
            'numero_nc': '001-001-0050486',
            'monto': Decimal('149348'),
            'factura_origen': '001-001-0288590',
            'fecha': '2026-10-27',
        },
        {
            'numero_nc': '001-001-0050485',
            'monto': Decimal('33024'),
            'factura_origen': '001-001-0288586',
            'fecha': '2026-10-27',
        }
    ]

    async with async_session_factory() as db:
        async with db.begin():
            print("\n1. Renombrando REM-20260926-004794 a factura 001-001-0280541...")
            await db.execute(text("""
                UPDATE supplier_invoices
                SET numero_factura = '001-001-0280541',
                    total = 552652,
                    saldo_pendiente = 552652,
                    notas = COALESCE(notas, '') || ' [Renombrado desde REM muelle a FAC oficial s/ planilla]'
                WHERE numero_factura = 'REM-20260926-004794' AND supplier_id = :sid
            """), {'sid': GONZALITO_ID})

            print("\n2. Calibrando valores nominales faciales de las 151 facturas del Excel...")
            calibradas_count = 0
            for inv_data in excel_invoices:
                num = inv_data['numero_factura']
                tot = inv_data['total_papel']

                res = await db.execute(text("""
                    SELECT id, total, saldo_pendiente
                    FROM supplier_invoices
                    WHERE supplier_id = :sid AND numero_factura = :num AND estado != 'cancelada'
                """), {'sid': GONZALITO_ID, 'num': num})
                row = res.fetchone()
                if row:
                    inv_id, cur_total, cur_saldo = row[0], row[1], row[2]
                    diff = tot - cur_total
                    if diff != 0:
                        calibradas_count += 1
                        await db.execute(text("""
                            UPDATE supplier_invoices
                            SET total = :tot,
                                estado = CASE WHEN estado = 'retenida_discrepancia' THEN 'aprobada' ELSE estado END,
                                notas = COALESCE(notas, '') || ' [Calibrado facial papel s/ Excel: dif=' || :diff || ']'
                            WHERE id = :id
                        """), {'tot': tot, 'diff': str(diff), 'id': inv_id})
                else:
                    print(f"  [!] Factura no encontrada en BD: {num}")

            print(f"   -> {calibradas_count} facturas calibradas a su monto impreso en papel exacto.")

            print("\n3. Registrando y vinculando las 57 Notas de Crédito válidas (55 Excel + 2 Nuevas)...")
            all_ncs_to_apply = excel_ncs + new_ncs_today
            valid_nc_nums = [nc['numero_nc'] for nc in all_ncs_to_apply]

            # Limpiar TODAS las aplicaciones de NCs anteriores de Casa Gonzalito
            await db.execute(text("""
                DELETE FROM supplier_credit_note_applications 
                WHERE invoice_id IN (SELECT id FROM supplier_invoices WHERE supplier_id = :sid)
            """), {'sid': GONZALITO_ID})

            # Cancelar cualquier otra NC de Gonzalito que no pertenezca a la verdad del Excel o las 2 nuevas
            await db.execute(text("""
                UPDATE supplier_credit_notes
                SET cancelado = true,
                    saldo_disponible = 0
                WHERE supplier_id = :sid AND NOT (numero = ANY(:valid_nums))
            """), {'sid': GONZALITO_ID, 'valid_nums': valid_nc_nums})

            nc_registradas = 0
            nc_aplicadas = 0

            for nc in all_ncs_to_apply:
                nc_num = nc['numero_nc']
                nc_amt = nc['monto']
                parent_fac = nc['factura_origen']
                nc_date = nc['fecha']

                # Obtener ID de factura padre
                fac_res = await db.execute(text("""
                    SELECT id FROM supplier_invoices
                    WHERE supplier_id = :sid AND numero_factura = :num AND estado != 'cancelada'
                """), {'sid': GONZALITO_ID, 'num': parent_fac})
                fac_row = fac_res.fetchone()
                if not fac_row:
                    print(f"  [!] Factura padre no encontrada para NC {nc_num}: {parent_fac}")
                    continue
                parent_id = fac_row[0]

                # Verificar si la NC ya existe
                nc_res = await db.execute(text("""
                    SELECT id FROM supplier_credit_notes
                    WHERE supplier_id = :sid AND numero = :num
                """), {'sid': GONZALITO_ID, 'num': nc_num})
                nc_row = nc_res.fetchone()

                # Parse date to datetime.date
                if isinstance(nc_date, datetime.date):
                    fec_obj = nc_date
                elif isinstance(nc_date, datetime.datetime):
                    fec_obj = nc_date.date()
                elif isinstance(nc_date, str):
                    fec_obj = datetime.date.fromisoformat(nc_date[:10])
                else:
                    fec_obj = datetime.date(2026, 10, 27)

                if not nc_row:
                    nc_id = uuid.uuid4()
                    await db.execute(text("""
                        INSERT INTO supplier_credit_notes (
                            id, company_id, supplier_id, numero, timbrado, fecha,
                            monto, saldo_disponible, cancelado, numero_factura_origen, motivo
                        ) VALUES (
                            :id, :cid, :sid, :num, :timb, :fec,
                            :monto, 0, false, :orig, 'Diferencia / Devolución s/ planilla liquidación'
                        )
                    """), {
                        'id': nc_id,
                        'cid': COMPANY_ID,
                        'sid': GONZALITO_ID,
                        'num': nc_num,
                        'timb': TIMBRADO_DEF,
                        'fec': fec_obj,
                        'monto': nc_amt,
                        'orig': parent_fac
                    })
                    nc_registradas += 1
                else:
                    nc_id = nc_row[0]
                    await db.execute(text("""
                        UPDATE supplier_credit_notes
                        SET monto = :monto,
                            saldo_disponible = 0,
                            cancelado = false,
                            numero_factura_origen = :orig
                        WHERE id = :id
                    """), {'monto': nc_amt, 'orig': parent_fac, 'id': nc_id})

                # Limpiar aplicaciones previas de esta NC
                await db.execute(text("DELETE FROM supplier_credit_note_applications WHERE credit_note_id = :ncid"), {'ncid': nc_id})

                # Insertar aplicación formal 1:1 a la factura padre
                await db.execute(text("""
                    INSERT INTO supplier_credit_note_applications (
                        id, company_id, credit_note_id, invoice_id, monto_aplicado, observaciones
                    ) VALUES (
                        gen_random_uuid(), :cid, :ncid, :invid, :monto, :obs
                    )
                """), {
                    'cid': COMPANY_ID,
                    'ncid': nc_id,
                    'invid': parent_id,
                    'monto': nc_amt,
                    'obs': f"Aplicación formal a FAC {parent_fac} s/ planilla liquidación Gonzalito"
                })
                nc_aplicadas += 1

            print(f"   -> {nc_registradas} nuevas NCs creadas en BD.")
            print(f"   -> {nc_aplicadas} NCs vinculadas formalmente a sus facturas.")

            print("\n4. Desactivando/Cancelando facturas espurias fuera del Excel...")
            # Cancelar 001-001-0261010 y 001-001-0282827
            for orphan_num in ['001-001-0261010', '001-001-0282827']:
                await db.execute(text("""
                    UPDATE supplier_invoices
                    SET estado = 'cancelada',
                        saldo_pendiente = 0,
                        notas = COALESCE(notas, '') || ' [Depurada/Cancelada s/ planilla oficial de corte]'
                    WHERE supplier_id = :sid AND numero_factura = :num AND saldo_pendiente > 0
                """), {'sid': GONZALITO_ID, 'num': orphan_num})
                print(f"   -> Factura {orphan_num} cancelada y saldo reseteado a 0.")

            print("\n5. Recalculando saldo pendiente contable inmutable de cada factura...")
            excel_doc_numbers = [inv['numero_factura'] for inv in excel_invoices]

            # 5.a Facturas del Excel (Deuda Real): saldo = total - sum(ncs_aplicadas)
            await db.execute(text("""
                UPDATE supplier_invoices inv
                SET saldo_pendiente = GREATEST(0, inv.total - COALESCE((
                        SELECT SUM(a.monto_aplicado)
                        FROM supplier_credit_note_applications a
                        WHERE a.invoice_id = inv.id
                    ), 0)),
                    estado = CASE 
                        WHEN (inv.total - COALESCE((SELECT SUM(a.monto_aplicado) FROM supplier_credit_note_applications a WHERE a.invoice_id = inv.id), 0)) <= 0 THEN 'pagada'
                        ELSE 'aprobada'
                    END
                WHERE inv.supplier_id = :sid AND inv.estado != 'cancelada' AND inv.numero_factura = ANY(:docs)
            """), {'sid': GONZALITO_ID, 'docs': excel_doc_numbers})

            # 5.b Facturas históricas fuera del Excel: saldadas en 0
            await db.execute(text("""
                UPDATE supplier_invoices inv
                SET saldo_pendiente = 0,
                    estado = 'pagada'
                WHERE inv.supplier_id = :sid 
                  AND inv.estado != 'cancelada'
                  AND NOT (inv.numero_factura = ANY(:docs))
            """), {'sid': GONZALITO_ID, 'docs': excel_doc_numbers})

        # Fin de la transacción atómica
        print("\n¡TRANSACCIÓN COMPLETADA EXITOSAMENTE!")

        # 6. Verificación posterior
        print("\n" + "=" * 80)
        print("VERIFICACIÓN DE CONCILIACIÓN FINAL:")
        res_fac = await db.execute(text("""
            SELECT COUNT(*) as count, SUM(total) as tot_facial, SUM(saldo_pendiente) as tot_saldo
            FROM supplier_invoices
            WHERE supplier_id = :sid AND estado != 'cancelada' AND saldo_pendiente > 0
        """), {'sid': GONZALITO_ID})
        r_fac = res_fac.fetchone()

        res_ncs = await db.execute(text("""
            SELECT COUNT(*) as count, SUM(saldo_disponible) as saldo_disp
            FROM supplier_credit_notes
            WHERE supplier_id = :sid AND cancelado = false
        """), {'sid': GONZALITO_ID})
        r_ncs = res_ncs.fetchone()

        print(f"Facturas activas con saldo:        {r_fac[0]}")
        print(f"Total Facial de Facturas con Saldo: {r_fac[1]:>15,.0f} Gs.")
        print(f"Saldo Pendiente Real Exigible:      {r_fac[2]:>15,.0f} Gs.")
        print(f"Saldo libre en monedero de NCs:     {r_ncs[1]:>15,.0f} Gs.")
        print("=" * 80)
        esperado_neto = Decimal('915997111') - Decimal('149348') - Decimal('33024')
        print(f"Esperado según Excel - Nuevas NCs:  {esperado_neto:>15,.0f} Gs.")
        diff_final = r_fac[2] - esperado_neto
        print(f"Diferencia Final:                   {diff_final:>15,.0f} Gs.")
        print("=" * 80)

asyncio.run(run_remediation())
