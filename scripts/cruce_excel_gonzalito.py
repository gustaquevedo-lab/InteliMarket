import openpyxl
from decimal import Decimal
import asyncio
from sqlalchemy import text
from api.src.db import async_session_factory

GONZALITO_ID = 'a2f37cf3-caa5-42b5-84cd-270a549d5a73'

async def compare():
    wb = openpyxl.load_workbook('datos/FACTURAS PENDIENTES GONZALITO.xlsx', data_only=True)
    sheet = wb.active
    
    excel_invs = {}
    excel_ncs = {}
    current_inv = None
    
    for i in range(1, sheet.max_row + 1):
        doc = sheet.cell(i, 2).value
        amt = sheet.cell(i, 3).value
        date_v = sheet.cell(i, 1).value
        if doc is None or amt is None:
            continue
        doc_str = str(doc).strip()
        amt_d = Decimal(str(amt))
        if amt_d > 0:
            current_inv = doc_str
            excel_invs[doc_str] = {'row': i, 'date': date_v, 'amt': amt_d, 'ncs': []}
        else:
            excel_ncs[doc_str] = {'row': i, 'date': date_v, 'amt': abs(amt_d), 'inv': current_inv}
            if current_inv and current_inv in excel_invs:
                excel_invs[current_inv]['ncs'].append({'doc': doc_str, 'amt': abs(amt_d)})
                
    print(f"Excel: {len(excel_invs)} facturas, {len(excel_ncs)} NCs")
    
    async with async_session_factory() as db:
        res = await db.execute(text('''
            SELECT id, numero_factura, fecha_emision, total, saldo_pendiente, estado
            FROM supplier_invoices
            WHERE supplier_id = :sid AND estado != 'cancelada' AND saldo_pendiente > 0
            ORDER BY numero_factura
        '''), {'sid': GONZALITO_ID})
        db_invs = {r.numero_factura.strip(): dict(r._mapping) for r in res.fetchall()}
        
        in_db_not_excel = set(db_invs.keys()) - set(excel_invs.keys())
        in_excel_not_db = set(excel_invs.keys()) - set(db_invs.keys())
        
        print(f"\n--- FACTURAS EN BD CON SALDO PERO NO EN EXCEL ({len(in_db_not_excel)}) ---")
        for f in sorted(in_db_not_excel):
            info = db_invs[f]
            print(f"  {f}: Total={info['total']:,.0f}, Saldo={info['saldo_pendiente']:,.0f}, Estado={info['estado']}")
            
        print(f"\n--- FACTURAS EN EXCEL PERO NO CON SALDO EN BD ({len(in_excel_not_db)}) ---")
        for f in sorted(in_excel_not_db):
            info = excel_invs[f]
            print(f"  {f}: Monto Excel={info['amt']:,.0f}")
            chk = await db.execute(text('''
                SELECT numero_factura, total, saldo_pendiente, estado
                FROM supplier_invoices
                WHERE supplier_id = :sid AND (numero_factura = :num OR numero_factura ILIKE :num_like)
            '''), {'sid': GONZALITO_ID, 'num': f, 'num_like': f"%{f[-7:]}%"})
            rows = chk.fetchall()
            if rows:
                for r in rows:
                    print(f"     -> Hallada en BD: {r.numero_factura} Total={r.total:,.0f} Saldo={r.saldo_pendiente:,.0f} Estado={r.estado}")
            else:
                print("     -> NO existe en BD con ningun estado")

        # Diferencias de montos
        matching = set(excel_invs.keys()) & set(db_invs.keys())
        diffs = []
        for f in matching:
            ex_amt = excel_invs[f]['amt']
            db_total = Decimal(str(db_invs[f]['total']))
            db_saldo = Decimal(str(db_invs[f]['saldo_pendiente']))
            diff = ex_amt - db_total
            if diff != 0:
                diffs.append((f, ex_amt, db_total, diff))
                
        print(f"\n--- FACTURAS COINCIDENTES: {len(matching)} | CON DIFERENCIA DE MONTO FACIAL: {len(diffs)} ---")
        diffs.sort(key=lambda x: abs(x[3]), reverse=True)
        for f, ex_amt, db_total, diff in diffs[:15]:
            print(f"  {f}: Excel={ex_amt:,.0f} | BD={db_total:,.0f} | Dif={diff:+,.0f}")
        total_facial_diff = sum(d[3] for d in diffs)
        print(f"Diferencia Neta Acumulada de Montos Faciales entre coincidentes: {total_facial_diff:+,.0f} Gs.")

        # Cruce de NCs
        print(f"\n--- CRUCE DE LAS 55 NOTAS DE CRÉDITO DEL EXCEL ---")
        res_ncs = await db.execute(text('''
            SELECT id, numero, timbrado, fecha, monto, saldo_disponible, cancelado, numero_factura_origen
            FROM supplier_credit_notes
            WHERE supplier_id = :sid
        '''), {'sid': GONZALITO_ID})
        db_ncs = {r.numero.strip(): dict(r._mapping) for r in res_ncs.fetchall()}

        nc_missing_in_db = []
        nc_amount_diff = []
        for nc_num, nc_data in excel_ncs.items():
            if nc_num not in db_ncs:
                # search partial
                matched = [k for k in db_ncs.keys() if nc_num[-7:] in k]
                if matched:
                    print(f"  NC Excel {nc_num} mapea a BD {matched[0]}")
                else:
                    nc_missing_in_db.append((nc_num, nc_data['amt'], nc_data['inv']))
            else:
                db_amt = Decimal(str(db_ncs[nc_num]['monto']))
                ex_amt = nc_data['amt']
                if db_amt != ex_amt:
                    nc_amount_diff.append((nc_num, ex_amt, db_amt, ex_amt - db_amt))

        print(f"NCs del Excel NO encontradas en BD: {len(nc_missing_in_db)}")
        for nc_num, amt, inv in nc_missing_in_db:
            print(f"  - NC {nc_num}: Monto={amt:,.0f} (afecta a FAC {inv})")
        print(f"NCs con diferencia de monto: {len(nc_amount_diff)}")
        for nc_num, ex_amt, db_amt, diff in nc_amount_diff:
            print(f"  - NC {nc_num}: Excel={ex_amt:,.0f} | BD={db_amt:,.0f} | Dif={diff:+,.0f}")

asyncio.run(compare())
